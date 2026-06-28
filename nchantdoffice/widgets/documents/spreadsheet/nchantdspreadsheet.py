# @@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@||
"""
---
<(META)>:
	docid:
	name: NchantdOffice Spreadsheet
	description: >
		NchantdOffice Spreadsheet Document Widget
	version: 0.0.1.0.0.0
	authority: filesystem
	security: seclvl2
	<(WT)>: -32
"""
# -*- coding: utf-8 -*-

# ======================================Standard Library Modules======================================================||
from os.path import dirname, join
import logging

# ======================================3rd Party Library Modules=====================================================||
from PyQt6.QtCore import Qt, pyqtSignal, QPoint
from PyQt6.QtGui import QColor, QFont
from PyQt6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QTableWidget, QTableWidgetItem,
    QHeaderView, QMenu, QToolBar, QPushButton, QLabel, QLineEdit,
    QFileDialog, QMessageBox, QTabWidget, QApplication
)

# ======================================Solutions Brewer Library Modules==============================================||
from kahndor import kahndor
from nchantrs.libraries import pyqt
from kahndor.logma import Logma
from nchantdoffice.widgets.documents.documents import NchantdOfficeDocument
from nchantdoffice.widgets.widgets import NchantdOfficeWidgetMixin

# ====================================================================================================================||
here = join(dirname(__file__), "")  # ||
logma = Logma(__name__)
logger = logging.getLogger(__name__)

# ====================================================================================================================||
pxcfg = join(here, "_data_", "spreadsheet.yaml")


class NchantdSpreadSheet(QWidget, NchantdOfficeWidgetMixin):
    """NchantdOffice Spreadsheet Document Widget"""
    
    VERSION = "0.0.1.0.0.0"
    
    # Signals
    dataChanged = pyqtSignal(dict)
    cellSelected = pyqtSignal(str, object)
    
    def __init__(self, cfg=None, parent=None):
        super().__init__(parent)
        self.cfg = cfg or {}
        self.parent_window = parent
        self.app = QApplication.instance()
        
        # Initialize config
        self.config = kahndor.Instruct(pxcfg)
        if parent:
            self.config.override(parent.config)
        self.config.override(cfg)
        
        # Spreadsheet state
        self.current_sheet = None
        self.selected_range = None
        self.clipboard = None
        self.is_workbook = self.cfg.get('is_workbook', False)
        
        # UI Components
        self.table = None
        self.toolbar = None
        self.formula_bar = None
        self.sheet_tabs = None
        
        self.setup_ui()
        self.setup_toolbar()
        self.connect_signals()
        
        # Load initial data if provided
        if cfg:
            self.load_from_config(cfg)
    
    def setup_ui(self):
        """Initialize the user interface"""
        main_layout = QVBoxLayout(self)
        main_layout.setContentsMargins(0, 0, 0, 0)
        main_layout.setSpacing(0)
        
        # Sheet tabs (for workbook mode)
        if self.is_workbook:
            self.sheet_tabs = QTabWidget()
            self.sheet_tabs.setTabsClosable(False)
            main_layout.addWidget(self.sheet_tabs)
        
        # Formula bar
        self.formula_bar = self.create_formula_bar()
        main_layout.addWidget(self.formula_bar)
        
        # Table widget
        self.table = QTableWidget()
        self.table.setAlternatingRowColors(True)
        self.table.setSelectionBehavior(QTableWidget.SelectionBehavior.SelectCells)
        self.table.setSelectionMode(QTableWidget.SelectionMode.ExtendedSelection)
        
        # Configure headers
        header = self.table.horizontalHeader()
        header.setSectionResizeMode(QHeaderView.ResizeMode.Interactive)
        header.setStretchLastSection(True)
        
        # Set default size
        self.table.setRowCount(100)
        self.table.setColumnCount(26)  # A-Z
        
        # Add column headers (A-Z)
        for i in range(26):
            self.table.setHorizontalHeaderItem(i, QTableWidgetItem(chr(65 + i)))
        
        # Add row headers (1-100)
        for i in range(100):
            self.table.setVerticalHeaderItem(i, QTableWidgetItem(str(i + 1)))
        
        main_layout.addWidget(self.table)
        self.setLayout(main_layout)
    
    def create_formula_bar(self):
        """Create the formula bar widget"""
        bar = QWidget()
        layout = QHBoxLayout(bar)
        layout.setContentsMargins(5, 5, 5, 5)
        
        # Cell reference label
        self.cell_ref_label = QLabel("A1")
        self.cell_ref_label.setMinimumWidth(60)
        layout.addWidget(QLabel("Cell:"))
        layout.addWidget(self.cell_ref_label)
        
        # Function button
        self.func_button = QPushButton("fx")
        self.func_button.setMaximumWidth(30)
        self.func_button.clicked.connect(self.show_function_dialog)
        layout.addWidget(self.func_button)
        
        # Formula input
        self.formula_input = QLineEdit()
        self.formula_input.setPlaceholderText("Enter value or formula (e.g., =SUM(A1:A10))")
        self.formula_input.returnPressed.connect(self.on_formula_entered)
        layout.addWidget(self.formula_input)
        
        layout.addStretch()
        return bar
    
    def setup_toolbar(self):
        """Create the toolbar"""
        self.toolbar = QToolBar("Spreadsheet Tools")
        
        # File operations
        self.toolbar.addAction("New", self.new_sheet)
        self.toolbar.addAction("Open", self.open_file)
        self.toolbar.addAction("Save", self.save_file)
        self.toolbar.addSeparator()
        
        # Edit operations
        self.toolbar.addAction("Cut", self.cut)
        self.toolbar.addAction("Copy", self.copy)
        self.toolbar.addAction("Paste", self.paste)
        self.toolbar.addSeparator()
        
        # Format operations
        self.toolbar.addAction("Bold", self.toggle_bold)
        self.toolbar.addAction("Italic", self.toggle_italic)
        self.toolbar.addSeparator()
        
        # Insert operations
        self.toolbar.addAction("Insert Row", self.insert_row)
        self.toolbar.addAction("Insert Column", self.insert_column)
        self.toolbar.addAction("Delete Row", self.delete_row)
        self.toolbar.addAction("Delete Column", self.delete_column)
        self.toolbar.addSeparator()
        
        # Chart insertion
        self.toolbar.addAction("Insert Chart", self.insert_chart)
    
    def connect_signals(self):
        """Connect UI signals to handlers"""
        self.table.itemChanged.connect(self.on_cell_changed)
        self.table.itemSelectionChanged.connect(self.on_selection_changed)
    
    # ================================ Data Loading ================================    
    
    def load_from_config(self, cfg):
        """Load spreadsheet data from configuration"""
        if 'data' in cfg:
            self.load_data(cfg['data'])
        if 'sheets' in cfg:
            for sheet_name, sheet_data in cfg['sheets'].items():
                self.add_sheet(sheet_name, sheet_data)
    
    def load_data(self, data):
        """Load data into the spreadsheet"""
        if isinstance(data, dict):
            # Load from cell dictionary
            for address, value in data.items():
                self.set_cell_value(address, value)
        elif isinstance(data, list):
            # Load from list (rows)
            for row_idx, row in enumerate(data):
                for col_idx, value in enumerate(row):
                    address = f"{chr(65 + col_idx)}{row_idx + 1}"
                    self.set_cell_value(address, value)
    
    def set_cell_value(self, address, value):
        """Set value at specified cell address (e.g., 'A1')"""
        if not address or len(address) < 2:
            return
        
        col = ord(address[0].upper()) - 65  # A=0, B=1, etc.
        row = int(address[1:]) - 1  # 1-indexed
        
        if 0 <= col < self.table.columnCount() and row >= 0:
            # Ensure row exists
            while row >= self.table.rowCount():
                self.table.insertRow(self.table.rowCount())
            
            item = self.table.item(row, col)
            if item is None:
                item = QTableWidgetItem(str(value))
                self.table.setItem(row, col, item)
            else:
                item.setText(str(value))
    
    def get_cell_value(self, address):
        """Get value from specified cell address"""
        if not address or len(address) < 2:
            return None
        
        col = ord(address[0].upper()) - 65
        row = int(address[1:]) - 1
        
        if 0 <= col < self.table.columnCount() and 0 <= row < self.table.rowCount():
            item = self.table.item(row, col)
            return item.text() if item else ""
        return None
    
    def get_current_address(self):
        """Get the currently selected cell address"""
        selected = self.table.selectedIndexes()
        if selected:
            row = selected[0].row()
            col = selected[0].column()
            return f"{chr(65 + col)}{row + 1}"
        return "A1"
    
    def get_selection_range(self):
        """Get the selected cell range"""
        selected = self.table.selectedIndexes()
        if not selected:
            return None, None
        
        rows = sorted(set(index.row() for index in selected))
        cols = sorted(set(index.column() for index in selected))
        
        if not rows or not cols:
            return None, None
        
        start_cell = f"{chr(65 + cols[0])}{rows[0] + 1}"
        end_cell = f"{chr(65 + cols[-1])}{rows[-1] + 1}"
        
        return start_cell, end_cell
    
    # ================================ Event Handlers ================================    
    
    def on_cell_changed(self, item):
        """Handle cell value change"""
        address = f"{chr(65 + item.column())}{item.row() + 1}"
        value = item.text()
        
        # Emit signal
        self.dataChanged.emit({'address': address, 'value': value})
    
    def on_selection_changed(self):
        """Handle selection change"""
        address = self.get_current_address()
        value = self.get_cell_value(address)
        
        # Update formula bar
        self.cell_ref_label.setText(address)
        self.formula_input.setText(value or "")
        
        # Emit signal
        self.cellSelected.emit(address, value)
    
    def on_formula_entered(self):
        """Handle formula input submission"""
        address = self.get_current_address()
        formula = self.formula_input.text()
        
        self.set_cell_value(address, formula)
    
    def show_function_dialog(self):
        """Show function selection dialog"""
        # TODO: Implement function dialog
        QMessageBox.information(self, "Functions", "Function dialog not yet implemented")
    
    # ================================ Toolbar Actions ================================    
    
    def new_sheet(self):
        """Create a new spreadsheet"""
        self.table.clearContents()
        self.dataChanged.emit({'action': 'new'})
    
    def open_file(self):
        """Open a spreadsheet file"""
        file_path, _ = QFileDialog.getOpenFileName(
            self,
            "Open Spreadsheet",
            "",
            "Spreadsheet Files (*.csv *.xlsx *.xls);;All Files (*)"
        )
        
        if file_path:
            self.load_from_file(file_path)
    
    def load_from_file(self, file_path):
        """Load spreadsheet from file"""
        try:
            if file_path.endswith('.csv'):
                self.load_csv(file_path)
            elif file_path.endswith('.xlsx') or file_path.endswith('.xls'):
                self.load_excel(file_path)
            else:
                QMessageBox.warning(self, "Error", "Unsupported file format")
        except Exception as e:
            QMessageBox.warning(self, "Error", f"Failed to load file: {e}")
    
    def load_csv(self, file_path):
        """Load CSV file into spreadsheet"""
        import csv
        
        self.table.clearContents()
        with open(file_path, 'r', encoding='utf-8') as f:
            reader = csv.reader(f)
            for row_idx, row in enumerate(reader):
                for col_idx, value in enumerate(row):
                    if col_idx >= self.table.columnCount():
                        self.table.insertColumn(col_idx)
                        self.table.setHorizontalHeaderItem(col_idx, QTableWidgetItem(chr(65 + col_idx)))
                    item = QTableWidgetItem(value)
                    self.table.setItem(row_idx, col_idx, item)
                
                if row_idx >= self.table.rowCount() - 1:
                    self.table.insertRow(row_idx + 1)
    
    def load_excel(self, file_path):
        """Load Excel file into spreadsheet"""
        try:
            import openpyxl
            wb = openpyxl.load_workbook(file_path)
            ws = wb.active
            
            self.table.clearContents()
            for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
                for col_idx, value in enumerate(row):
                    if col_idx >= self.table.columnCount():
                        self.table.insertColumn(col_idx)
                        self.table.setHorizontalHeaderItem(col_idx, QTableWidgetItem(chr(65 + col_idx)))
                    item = QTableWidgetItem(str(value) if value is not None else "")
                    self.table.setItem(row_idx, col_idx, item)
                
                if row_idx >= self.table.rowCount() - 1:
                    self.table.insertRow(row_idx + 1)
        except ImportError:
            QMessageBox.warning(self, "Error", "openpyxl not installed. Please install: pip install openpyxl")
        except Exception as e:
            QMessageBox.warning(self, "Error", f"Failed to load Excel: {e}")
    
    def save_file(self):
        """Save spreadsheet to file"""
        file_path, _ = QFileDialog.getSaveFileName(
            self,
            "Save Spreadsheet",
            "",
            "CSV Files (*.csv);;Excel Files (*.xlsx);;All Files (*)"
        )
        
        if file_path:
            self.save_to_file(file_path)
    
    def save_to_file(self, file_path):
        """Save spreadsheet to file"""
        try:
            if file_path.endswith('.csv') or not file_path.endswith('.xlsx'):
                self.save_csv(file_path if file_path.endswith('.csv') else file_path + '.csv')
            elif file_path.endswith('.xlsx'):
                self.save_excel(file_path)
            
            QMessageBox.information(self, "Success", f"File saved: {file_path}")
        except Exception as e:
            QMessageBox.warning(self, "Error", f"Failed to save file: {e}")
    
    def save_csv(self, file_path):
        """Save spreadsheet to CSV file"""
        import csv
        
        with open(file_path, 'w', encoding='utf-8', newline='') as f:
            writer = csv.writer(f)
            for row in range(self.table.rowCount()):
                row_data = []
                has_data = False
                for col in range(self.table.columnCount()):
                    item = self.table.item(row, col)
                    value = item.text() if item else ""
                    row_data.append(value)
                    if value:
                        has_data = True
                if has_data:
                    writer.writerow(row_data)
    
    def save_excel(self, file_path):
        """Save spreadsheet to Excel file"""
        try:
            import openpyxl
            wb = openpyxl.Workbook()
            ws = wb.active
            
            for row in range(self.table.rowCount()):
                for col in range(self.table.columnCount()):
                    item = self.table.item(row, col)
                    value = item.text() if item else ""
                    if value:
                        ws.cell(row=row + 1, column=col + 1, value=value)
            
            wb.save(file_path)
        except ImportError:
            QMessageBox.warning(self, "Error", "openpyxl not installed")
    
    # ================================ Clipboard Operations ================================    
    
    def cut(self):
        """Cut selected cells"""
        self.copy()
        self.clear_selection()
    
    def copy(self):
        """Copy selected cells"""
        selected = self.table.selectedIndexes()
        if not selected:
            return
        
        # Get selection range
        rows = sorted(set(index.row() for index in selected))
        cols = sorted(set(index.column() for index in selected))
        
        # Copy to clipboard
        clipboard_data = []
        for row in rows:
            row_data = []
            for col in cols:
                item = self.table.item(row, col)
                row_data.append(item.text() if item else "")
            clipboard_data.append("\t".join(row_data))
        
        self.clipboard = "\n".join(clipboard_data)
        clipboard = QApplication.clipboard()
        clipboard.setText(self.clipboard)
    
    def paste(self):
        """Paste from clipboard"""
        clipboard = QApplication.clipboard()
        text = clipboard.text()
        
        if not text:
            return
        
        selected = self.table.selectedIndexes()
        if not selected:
            return
        
        start_row = selected[0].row()
        start_col = selected[0].column()
        
        # Parse clipboard data
        rows = text.split("\n")
        for row_idx, row_data in enumerate(rows):
            cols = row_data.split("\t")
            for col_idx, value in enumerate(cols):
                target_row = start_row + row_idx
                target_col = start_col + col_idx
                
                # Ensure rows exist
                while target_row >= self.table.rowCount():
                    self.table.insertRow(self.table.rowCount())
                
                # Ensure columns exist
                while target_col >= self.table.columnCount():
                    self.table.insertColumn(self.table.columnCount())
                    self.table.setHorizontalHeaderItem